import fitz
import re
import openai
import time
import logging
import json
from docx import Document 
from pdf2image import convert_from_bytes
from pdf2image import convert_from_path
import os
from docx import Document
# from IPython.display import Image, display
import base64
import openai
import shutil
import cv2
import numpy as np
from PyPDF2 import PdfReader, PdfWriter
from PIL import Image
import pytesseract
import streamlit as st
from io import BytesIO
from langchain.chat_models import ChatOpenAI
from langchain.prompts import PromptTemplate
from langchain_core.output_parsers import JsonOutputParser
pytesseract.pytesseract.tesseract_cmd = r'Tesseract-OCR\tesseract.exe'

# Load the config file
with open('Config/configuration.json', 'r') as f:
    config = json.load(f)

openai.api_key = r"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"
config['api_key']

logging.basicConfig(
    filename='logs/app.log',
    filemode='a',  # Append mode
    format='%(asctime)s - %(levelname)s - %(message)s',
    level=logging.INFO
)

image_list=[]
def extract_pdf_text(pdf_file, page_number=[]):
    """
    Extract text from the PDF while attempting to include all content, including tables.
    """
    extracted_text = []
    try:
        # Open the PDF from bytes
        doc = fitz.open(stream=pdf_file, filetype="pdf")
        
        if len(page_number) == 0:
            # Loop over the first 45 pages or up to total page count
            for page_num in range(min(45, doc.page_count)):
                try:
                    page = doc.load_page(page_num)
                    blocks = page.get_text("blocks")  # Extract text blocks with positions

                    for block in blocks:
                        block_text = block[4]  # Text content of the block
                        extracted_text.append(block_text)
                except Exception as e:
                    st.error(f"Failed to load page {page_num}: {e}")
        else:
            for page_num in page_number:
                try:
                    page = doc.load_page(page_num)
                    blocks = page.get_text("blocks")  # Extract text blocks with positions

                    for block in blocks:
                        block_text = block[4]  # Text content of the block
                        extracted_text.append(block_text)
                except Exception as e:
                    st.error(f"Failed to load page {page_num}: {e}")

    finally:
        doc.close()  # Ensure the document is closed even if an error occurs

    return "\n".join(extracted_text)

def extract_text_from_word(docx_file):
    """
    Extracts text from a Word document without any additional cleaning or processing.
    """
    doc = Document(docx_file)
    all_text = []
    for paragraph in doc.paragraphs:
        all_text.append(paragraph.text)
        pdf_content = "\n".join(all_text)
    return pdf_content

# convert pdf page to images
def pdf_to_images(pdf_bytes):
    """
    Converts a PDF file (in bytes) to a list of images, one for each page.
    """
    try:
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        images = []
        for page_num in range(len(doc)):
            page = doc.load_page(page_num)
            pix = page.get_pixmap()
            image = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)
            images.append(image)
        return images
    except Exception as e:
        logging.error(f"Error converting PDF to images: {e}")
        return []

# yellow color detection in page
def contains_yellow(image):
    """
    Checks if the given image contains yellow color.
    """
    try:
        hsv_image = cv2.cvtColor(image, cv2.COLOR_RGB2HSV)
        yellow_mask = cv2.inRange(hsv_image, (20, 100, 100), (30, 255, 255))
        return np.any(yellow_mask)
    except cv2.error as e:
        logging.error(f"OpenCV error during color conversion or masking: {e}")
        return False
    except Exception as e:
        logging.error(f"Unexpected error in contains_yellow function: {e}")
        return False

# find yellow color and save in folder
def extract_images_and_figures_page_number(pdf_path,output_folder):
    global image_list
    images =  pdf_to_images(pdf_path)
    
    pdf_stream = BytesIO(pdf_path)
    reader = PdfReader(pdf_stream)
    writer = PdfWriter()
    pages_chosen = []
    
    for page_num, image in enumerate(images):
        page = reader.pages[page_num]
        text = page.extract_text().lower()
        
        keywords=["hazard","warning","caution","Precaution"]
        if contains_yellow(image) or any(word in text.lower() for word in keywords):
            writer.add_page(page)
            pages_chosen.append(page_num+1)

    if os.path.exists(output_folder):
        shutil.rmtree(output_folder)
    os.makedirs(output_folder, exist_ok=True)
    
    # image_list = convert_from_path(pdf_path)    
    for page in pages_chosen:
        # PDF pages are 1-indexed
        images_from_page = convert_from_bytes(pdf_path, first_page=page, last_page=page)
    
        image_path = f"{output_folder}/page_{page}.png"
        images_from_page[0].save(image_path, 'JPEG')

# encode the image to base64 
def encode_image(image_path):
    try:
        with open(image_path, "rb") as image_file:
            return base64.b64encode(image_file.read()).decode("utf-8")
    except IOError as e:
        logging.error(f"Error opening image file {image_path}: {e}")
        return None
    
def image_based_warning(folder_name):
    image_paths=[]
    image_name=[]
    for f in os.listdir(folder_name):
        image_name.append(f)
        image_paths.append(f"{folder_name}\\{f}")
        # ]
    image_text=[]
    input_tokens=0
    output_tokens=0
    total_tokens=0
    for idx, image_path in enumerate(image_paths, start=0):
            messages = [
                {
            "role": "system",
            "content": "Analyze the following text and extract all **WARNINGS,NOTE,IMPORTANT NOTE and CAUTIONS**. If some warnings or cautions do not **explicitly use these keywords but imply risks, hazards, or safety precautions, extract those all statements as well**. Focus only on the text that includes a yellow triangle symbol on the left-hand side or any hazard-related symbols. **Note:- 1)Exclude any predefined template sentences that describe what symbols indicate, such as warnings about electrical shock, fire, sharp points, hot surfaces, gloves, or pinch points. 2) Do not add any extra information or comments—respond only with the extracted text.\n3) If no relevant content is found, return an empty response without any comment or apology**. Ensure the extraction is precise and captures all risk-related information associated with those symbols."
            },]
            base64_image = encode_image(image_path)
            messages.append({
                "role": "user",
                "content": [
                    {"type": "text", "text": "Analyze the following text and extract all **WARNINGS,NOTE,IMPORTANT NOTE and CAUTIONS**. If some warnings or cautions do not explicitly use these keywords but imply risks, hazards, or safety precautions, extract those statements as well. extract and focus only on the text that includes a yellow triangle symbol on the left-hand side or any hazard-related symbols. **Note:- Exclude any predefined template sentences that describe what symbols indicate, such as warnings about electrical shock, fire, sharp points, hot surfaces, gloves, or pinch points** and **Do not add any extra information or comments—respond only with the extracted text **, Ensure the extraction is precise and captures all risk-related information associated with those symbols."},
                    {"type": "image_url", "image_url": {
                        "url": f"data:image/png;base64,{base64_image}"}
                    }
                ]
            })

            start_time=time.time()
            response = openai.ChatCompletion.create(
                    model=config["model_name"],  # Replace with your model
                    messages=messages,
                    temperature=0.0,
                )  
            response_time = time.time() - start_time

            if response and response.choices and response.choices[0].message.content:
                    input_tokens += response['usage']['prompt_tokens']
                    output_tokens += response['usage']['completion_tokens']
                    total_tokens += response['usage']['total_tokens']
                    image_text.append(response.choices[0].message.content)

    logging.info(f"image_based_warning Section.................")
    logging.info(f"Input Tokens: {input_tokens}")
    logging.info(f"Output Tokens: {output_tokens}")
    logging.info(f"Total Tokens: {total_tokens}")
    logging.info(f"Response generation time: {response_time:.2f} seconds")

                    
    return image_text,{
            'input_tokens': input_tokens,
            'output_tokens': output_tokens,
            'total_tokens': total_tokens,
            'response_time': response_time
        }

def process_text_with_GPT(text, reference_content):
    """
    Sends the extracted text to the OpenAI API for processing based on the reference content provided.
    """
    prompt = f"""You must refer strictly to the reference template provided below and perform the following tasks:
    
    - **Do not include** any acknowledgments of the model's capabilities or limitations.
    - Focus exclusively on the **"1. Contraindications**" and "**2. Warnings, Precautions and Potential Adverse Effects"** in the 'TEXT Extraction' section within the reference template,while excluding any ** 'Warnings and Precautions' ** or'table' or 'image' sections.
    - Do not repeat any sections from the reference template or add any additional information.
    - Do not print any conversational responses generated by the model.
    - **also write heading bold 1.) **Contraindications"** and "2. **Warnings, Precautions and Potential Adverse Effects"**

Reference Template:\n{reference_content}\n
Content to Process:\n{text}
"""

    start_time = time.time()

    response = openai.ChatCompletion.create(
        model=config['model_name'],  # Use GPT-4 or GPT-4-turbo based on your configuration
        messages=[
            {"role": "system", "content": "You are a helpful assistant."},
            {"role": "user", "content": prompt}
        ],
        temperature=config['generation_config']['temperature'],
        max_tokens=config['generation_config']['max_tokens'],
        top_p=config['generation_config']['top_p'],
    )

    response_time = time.time() - start_time

    # Return the generated content if available
    if response and response['choices']:
        input_tokens = response['usage']['prompt_tokens']
        output_tokens = response['usage']['completion_tokens']
        total_tokens = response['usage']['total_tokens']

    # Log the token usage and response time
        logging.info("process_text_with_GPT....")
        logging.info(f"Input Tokens: {input_tokens}")
        logging.info(f"Output Tokens: {output_tokens}")
        logging.info(f"Total Tokens: {total_tokens}")
        logging.info(f"Response generation time: {response_time:.2f} seconds")
        return response['choices'][0]['message']['content'], {
            'input_tokens': input_tokens,
            'output_tokens': output_tokens,
            'total_tokens': total_tokens,
            'response_time': response_time
        }
    else:
        return None
def process_warning_text_with_GPT(text, reference_content):
    """
    Sends the extracted text to the OpenAI API for processing based on the reference content provided.
    """
    prompt = f"""This is the "context": {text}.
        **write heading on the top Warning and Precaution **  
        Your task is to extract **all information** from the provided list, focusing solely on meaningful content, including warnings, cautions, important note and instructions.  

        **Instructions for Extraction:**  
        1. **Extract all warnings and cautions** (whether explicitly labeled or implied) and any other relevant safety instructions.  
        2. **Exclude**:
        - **Symbol indications** (e.g., ⚠️, hazard triangles).
        - **Model or system capability-related statements**, such as: "I'm sorry, I can't assist with that."
        3. **Do not extract any statments from the legend which contains the word "This symbol indicates".**
        4. must**Do not include any acknowledgments of the model's capabilities or limitations**
        5. **Maintain the original structure** of the text and return **each item in its complete form**.  
        6. **Do not alter, summarize, or omit** any part of the extracted content. Keep all relevant warnings, cautions, and instructions intact.
        
        
        Focus on clear and comprehensive extraction following the above rules.

        
        Reference Template:\n{reference_content}\n
       
        """

    start_time = time.time()

    response = openai.ChatCompletion.create(
        model=config['model_name'],  # Use GPT-4 or GPT-4-turbo based on your configuration
        messages=[
            {"role": "system", "content": "You are a helpful assistant."},
            {"role": "user", "content": prompt}
        ],
        temperature=config['generation_config']['temperature'],
        max_tokens=config['generation_config']['max_tokens'],
        top_p=config['generation_config']['top_p'],
    )

    response_time = time.time() - start_time
    
    
    # Return the generated content if available
    if response and response['choices']:
        input_tokens = response['usage']['prompt_tokens']
        output_tokens = response['usage']['completion_tokens']
        total_tokens = response['usage']['total_tokens']

    # Log the token usage and response time
        logging.info("process_warning_text_with_GPT...............")
        logging.info(f"Input Tokens: {input_tokens}")
        logging.info(f"Output Tokens: {output_tokens}")
        logging.info(f"Total Tokens: {total_tokens}")
        logging.info(f"Response generation time: {response_time:.2f} seconds")
        return response['choices'][0]['message']['content'], {
            'input_tokens': input_tokens,
            'output_tokens': output_tokens,
            'total_tokens': total_tokens,
            'response_time': response_time
        }
    else:
        return None, None

def image_discription_text(pdf_path, index):
    all_text = ''
    for page in index:
        try:
            images_from_page = convert_from_bytes(pdf_path, first_page=page, last_page=page)
            text = pytesseract.image_to_string(images_from_page[0])
            all_text += "/n/n" + text
        except Exception as e:
            logging.error(f"Error processing page {page}: {e}")
    return all_text

def process_device_discription_with_GPT_cer(text, reference_content):
    """
    Sends the extracted text to the OpenAI API for processing based on the reference content provided.
    """
    prompt = f"""You must refer strictly to the Reference Template provided below and perform the following tasks:
    - **Do not include** any acknowledgments of the model's capabilities or limitations.
    - 
    - **focus only "Device Description" section of the Reference Template**'. do not extract any other section information
    - **Identify each section through the headings in the Reference Template and extract all information from PDF Text, do not modify or trim any content.
Reference Template:\n{reference_content}\n
PDF Text:\n{text}
"""
    start_time = time.time()

    response = openai.ChatCompletion.create(
        model=config['model_name'],  # Use GPT-4 or GPT-4-turbo based on your configuration
        messages=[
            {"role": "system", "content": "You are a helpful assistant."},
            {"role": "user", "content": prompt}
        ],
        temperature=config['generation_config']['temperature'],
        max_tokens=config['generation_config']['max_tokens'],
        top_p=config['generation_config']['top_p'],
    )

    response_time = time.time() - start_time
    # Return the generated content if available
    if response and response['choices']:
         input_tokens = response['usage']['prompt_tokens']
         output_tokens = response['usage']['completion_tokens']
         total_tokens = response['usage']['total_tokens']

    # Log the token usage and response time
         logging.info("process_device_discription_with_GPT_cer...............")
         logging.info(f"Input Tokens: {input_tokens}")
         logging.info(f"Output Tokens: {output_tokens}")
         logging.info(f"Total Tokens: {total_tokens}")
         logging.info(f"Response generation time: {response_time:.2f} seconds")
         return response['choices'][0]['message']['content'], {
            'input_tokens': input_tokens,
            'output_tokens': output_tokens,
            'total_tokens': total_tokens,
            'response_time': response_time
        }

    else:
        return None
    
def process_text_with_GPT_cer(text, reference_content,indication):
    """
    Sends the extracted text to the OpenAI API for processing based on the reference content provided.
    """
    prompt = f"""You must refer strictly to the Reference Template provided below and perform the following tasks:
    - **Do not include** any acknowledgments of the model's capabilities or limitations.
    - **Use the reference template to receive your instructions and to provide a structure to the output, do not extract any content from the reference template**
    -**Ensure that all headings and sub-headings are printed in bold**
    - **Focus exclusively on all sections that come after the "Table Extraction" section of the Reference Template. Do not include information from any sections preceding "Table Extraction." or"Table Extraction**.
    - **Identify each section through the headings in the Reference Template and extract all information from PDF Text, do not modify or trim any content.
     Note:- The following conditions apply only for the 'Similar Devices' and 'Equivalent Devices' sections in the reference template:

     If Indication is 0, then print the generic name of the device in the 'Similar Devices' section provided in the reference template and print the following in the 'Equivalent Devices' section provided in the template: 
     “Thermo Fisher Scientific (Asheville) LLC has elected not to use the clinical data from an equivalent (clinical, technical and biological characteristics) device(s). 
     In the event, there are devices considered equivalent, their data will be considered as similar devices.”
    If Indication is 1, print "User input required" in both the sections. 
     
        Reference Template:\n{reference_content}\n
        PDF Text:\n{text}
        indication:\n{indication}\n

        """

    start_time = time.time()

    response = openai.ChatCompletion.create(
        model=config['model_name'],  # Use GPT-4 or GPT-4-turbo based on your configuration
        messages=[
            {"role": "system", "content": "You are a helpful assistant."},
            {"role": "user", "content": prompt}
        ],
        temperature=config['generation_config']['temperature'],
        max_tokens=config['generation_config']['max_tokens'],
        top_p=config['generation_config']['top_p'],
    )

    response_time = time.time() - start_time

    # Return the generated content if available
    if response and response['choices']:
        input_tokens = response['usage']['prompt_tokens']
        output_tokens = response['usage']['completion_tokens']
        total_tokens = response['usage']['total_tokens']

    # Log the token usage and response time
        logging.info("process_text_with_GPT_cer...............")
        logging.info(f"Input Tokens: {input_tokens}")
        logging.info(f"Output Tokens: {output_tokens}")
        logging.info(f"Total Tokens: {total_tokens}")
        logging.info(f"Response generation time: {response_time:.2f} seconds")
        return response['choices'][0]['message']['content'], {
            'input_tokens': input_tokens,
            'output_tokens': output_tokens,
            'total_tokens': total_tokens,
            'response_time': response_time
        }
    
    else:
        return None, None
    
############################### dynamic Tenplate #########################################################

def key_stucture(json_text):
    merged_data = {}
    previous_key=""
    current_key = None
    current_value = ""
    for key, value in json_text.items():
        # Extract base key by removing digits and hyphens
        base_key = re.sub(r'-\d+', '', key)
        
        if current_key is None:  # Start with the first key
            
            previous_key=key
            current_key = base_key
            current_value = value
            
        elif base_key == current_key:  # Continue merging if the key matches
            current_value += f"\n\n{value}"
        else:  # Different key encountered, save the current group and start a new one
            merged_data[previous_key] = current_value
            current_key = base_key
            current_value = str(value)
            previous_key=key

    # Add the last merged group to the dictionary
    if current_key:
        merged_data[key] = current_value
        
    return merged_data

def dynamic_reference(pdf_content):
    start_time = time.time()
    response = openai.ChatCompletion.create(
        model="gpt-4o",
        messages=[
            {  
                "role": "system",
                "content": "You are a helpful assistant that outputs only JSON. **Do not include any additional text or comments outside of the JSON**.** And do not ommiting any heading Number like 1,2,3.... and Instuction or Note**"
            },
            {   
                "role": "user",
                "content": f"""
                You are a helpful assistant that outputs **only JSON** without any additional text or comments outside the JSON. adhering to the following rules:
                    
                    1. **Preserve All Headings with Numbers**: Every heading in the input document must appear exactly as it is, including its number. Do not modify or omit the numbers in the headings.

                    2. **Category Differentiation**:
                    - The document is divided into three main categories: `Text`, `Table`, and `Image`.
                    - If any category appears multiple times in the input, append a unique identifier (e.g., `Text-1`, `Text-2`, `Table-1`, etc.) to differentiate them.
                    - Use the main category names (`Text`, `Table`, `Image`) as top-level keys, and append identifiers only when necessary.
                    - **don't make sub key in main key write as its in key value**
            
                    3. **Retain Original Formatting**:
                    - For sections labeled as `Text`, retain the input content **as-is do not ommiting any text**, including any numbers, labels, or special formatting.
                    - **For sections labeled as `Table`, use the given headings and structure provided in the input. **please Do not add placeholders like "user input required" beside column names****.

                    4. **Handle Numbers with Labels**:
                    - If the document contains numbers with labels (e.g., "1 Introduction", "Table 3 Device Features"), retain them exactly as they are in the output. Do not alter or remove these numbers.
                    **Don't add "user input required" in table section infront of columns list **
                   
                        
                    Here is the document content:

                    DOC Text: {pdf_content}

                    Generate the structured JSON output strictly adhering to the above instructions.


                """
            }
        ],
        temperature=0.0,
        response_format={"type": "json_object"}
    )
    
    response_time = time.time() - start_time
    if response and response['choices']:
        input_tokens = response['usage']['prompt_tokens']
        output_tokens = response['usage']['completion_tokens']
        total_tokens = response['usage']['total_tokens']

        # Log the token usage and response time
        logging.info("dynamic_reference.................")
        logging.info(f"Input Tokens: {input_tokens}")
        logging.info(f"Output Tokens: {output_tokens}")
        logging.info(f"Total Tokens: {total_tokens}")
        logging.info(f"Response generation time: {response_time:.2f} seconds")
    
    # Parse the JSON response from the LLM
    response_text = response['choices'][0]['message']['content'].strip()
    
    json_output = json.loads(response_text)
    json_output1=key_stucture(json_output)
    return json_output1,{"input_tokens": input_tokens, "output_tokens": output_tokens, "total_tokens": total_tokens, "response_time": response_time}

def process_text_with_GPT_cer_1(text, reference_content,indication=0):
    start_time = time.time()
    print("-----------")
    print(reference_content)
    print("-----------")

    """
    Sends the extracted text to the OpenAI API for processing based on the reference content provided.
    """
    model = ChatOpenAI(model="gpt-4o",api_key=config['api_key'],temperature=0.0, verbose=True)  # Enable verbose mode for token logging
    prompt_template = """
        You must strictly adhere to the Reference Template provided below and perform the following tasks:
       instruction:-
        -Extract values based on the structure outlined in the Reference Template.
        -Do not include any comments or acknowledgments about the model's capabilities or limitations.
        -Follow the Reference Template as your sole guide to structure the output. **Do not extract any content directly from the Reference Template itself**.
        -**Ensure that all headings and subheadings are displayed in bold for clarity**.
        -**If any headings have number in starting,retain the numbering in the output headings else give the headings as it is without the numbers**
        -Identify each section using the headings from the Reference Template and extract all relevant information from the PDF Text. Do not modify, trim, or **omit any content during extraction** 
        -**Do not mentions or generate any Class, classified or classification of device in any section like class B,Class II, etc**.
        
        Inputs:-

        Reference Template:
        {reference_content}

        PDF Text:
        {text}

        Indication:
        {indication}
        """
    prompt = PromptTemplate(template=prompt_template, input_variables=["text", "reference_content", "indication"])

    # Step 3: Set Up the Output Parser and Chain
    json_parser = JsonOutputParser()
    chain = prompt | model 

    # Desired columns and model number
    


    extracted_data = chain.invoke({"text": text, "reference_content": reference_content, "indication": indication})
    
    response_time = time.time() - start_time

    # Log the token usage and response time if available
    input_tokens = output_tokens = total_tokens = 0
    chain_tokens = {}  # This will store the chain token details
    if hasattr(extracted_data, 'usage'):
        input_tokens = extracted_data.usage.get('prompt_tokens', 0)
        output_tokens = extracted_data.usage.get('completion_tokens', 0)
        total_tokens = extracted_data.usage.get('total_tokens', 0)
        
        chain_tokens = {
            "input_tokens": input_tokens,
            "output_tokens": output_tokens,
            "total_tokens": total_tokens
        }

        logging.info("process_text_with_GPT_cer_1.................")
        logging.info(f"Input Tokens: {input_tokens}")
        logging.info(f"Output Tokens: {output_tokens}")
        logging.info(f"Total Tokens: {total_tokens}")
        logging.info(f"Response generation time: {response_time:.2f} seconds")

    # Return the extracted content and chain token details
    return extracted_data.content, chain_tokens
